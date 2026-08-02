from __future__ import annotations

import csv
import re
from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path

from bs4 import BeautifulSoup
from bs4.element import Comment, NavigableString, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "copy-feedback"
XLSX_PATH = OUT_DIR / "FCG Website Copy Review.xlsx"
CSV_PATH = OUT_DIR / "FCG Website Copy Review.csv"

PAGES = [
    "index.html",
    "about.html",
    "atlas.html",
    "agreeease.html",
    "api-access.html",
    "ai-booking-tool.html",
    "agent-booking-tool.html",
    "corporate-booking-tool.html",
    "ai-mapping.html",
    "whitelabel.html",
    "advisory-implementation.html",
    "index-variation-1.html",
]

HEADERS = [
    "Page filename",
    "Approx. section/context",
    "Element / tag / attribute source",
    "Text string",
    "Character count",
    "Suggested feedback owner",
    "Status",
    "Teammate comments",
    "Priority",
    "Extraction note",
]

IGNORED_TAGS = {
    "script",
    "style",
    "svg",
    "path",
    "defs",
    "clipPath",
    "template",
    "canvas",
}

GROUPING_TAGS = {
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "p",
    "a",
    "button",
    "label",
    "legend",
    "summary",
    "figcaption",
    "blockquote",
    "li",
    "td",
    "th",
    "option",
}

MINOR_TAGS = {
    "span",
    "div",
    "strong",
    "em",
    "small",
    "sup",
    "sub",
    "dt",
    "dd",
}

USEFUL_ATTRIBUTES = ("alt", "aria-label", "title")


def clean_text(value: str) -> str:
    value = value.replace("\xa0", " ")
    value = re.sub(r"[ \t\r\f\v]+", " ", value)
    value = re.sub(r"\n\s*", "\n", value)
    return value.strip()


def fingerprint(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def is_code_like(value: str) -> bool:
    if not value:
        return True
    if re.fullmatch(r"[\W_]+", value):
        return True
    if len(value) > 5000:
        return True
    if re.search(r"(function\s*\(|const\s+|let\s+|var\s+|=>|</?[a-z][^>]*>)", value):
        return True
    if len(value) > 80 and re.fullmatch(r"[A-Za-z0-9_\-:/.[\]#%?=&, ]+", value):
        if sum(ch.isupper() for ch in value) > 20 and " " not in value[:25]:
            return True
    return False


def is_generic_attribute(value: str) -> bool:
    return fingerprint(value) in {
        "",
        "placeholder",
        "logo",
        "image",
        "icon",
        "decorative",
    }


def has_ancestor(tag: Tag, names: set[str]) -> bool:
    parent = tag.parent
    while isinstance(parent, Tag):
        if parent.name in names:
            return True
        parent = parent.parent
    return False


def inside_ignored(tag: Tag | NavigableString) -> bool:
    parent = tag.parent if isinstance(tag, NavigableString) else tag
    while isinstance(parent, Tag):
        if parent.name in IGNORED_TAGS:
            return True
        parent = parent.parent
    return False


def section_context(tag: Tag | None) -> str:
    if tag is None:
        return "Head metadata"

    current = tag
    best: Tag | None = None
    while isinstance(current, Tag):
        if current.name == "nav":
            return "Navigation (shared)"
        if current.name == "footer":
            return "Footer (shared)"
        if current.name in {"section", "article", "header", "main"}:
            best = current
            break
        current = current.parent

    if best is None:
        return "Body copy"

    label = []
    if best.name:
        label.append(best.name)
    if best.get("id"):
        label.append(f"#{best.get('id')}")

    classes = best.get("class") or []
    helpful_classes = [
        c
        for c in classes
        if any(token in c.lower() for token in ["hero", "feature", "showcase", "bento", "section", "footer", "nav", "card"])
    ]
    if helpful_classes and not best.get("id"):
        label.extend(f".{c}" for c in helpful_classes[:2])

    heading = best.find(re.compile("^h[1-6]$"))
    if heading:
        heading_text = clean_text(heading.get_text(" ", strip=True))
        if heading_text:
            label.append(f"- {heading_text[:80]}")

    return " ".join(label) if label else "Body copy"


def add_row(rows: list[dict[str, object]], page: str, context: str, source: str, text: str, note: str = "") -> None:
    text = clean_text(text)
    if is_code_like(text):
        return
    rows.append(
        {
            "Page filename": page,
            "Approx. section/context": context,
            "Element / tag / attribute source": source,
            "Text string": text,
            "Character count": len(text),
            "Suggested feedback owner": "",
            "Status": "",
            "Teammate comments": "",
            "Priority": "",
            "Extraction note": note,
        }
    )


def extract_page(page: str) -> list[dict[str, object]]:
    html = (ROOT / page).read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "lxml")
    rows: list[dict[str, object]] = []

    title = soup.find("title")
    if title:
        add_row(rows, page, "Head metadata", "title", title.get_text(" ", strip=True))

    for meta in soup.find_all("meta"):
        name = (meta.get("name") or meta.get("property") or "").lower()
        content = meta.get("content") or ""
        if name in {"description", "og:description", "twitter:description", "og:title", "twitter:title"}:
            add_row(rows, page, "Head metadata", f"meta[{name}]", content)

    body = soup.body
    if not body:
        return rows

    for element in body.find_all(True):
        if element.name in IGNORED_TAGS or inside_ignored(element):
            continue

        for attr in USEFUL_ATTRIBUTES:
            value = element.get(attr)
            if isinstance(value, list):
                value = " ".join(value)
            if value and not is_generic_attribute(value):
                add_row(
                    rows,
                    page,
                    section_context(element),
                    f"{element.name}[{attr}]",
                    str(value),
                    "UX/accessibility attribute",
                )

        if element.name in GROUPING_TAGS:
            if has_ancestor(element, GROUPING_TAGS):
                continue
            add_row(rows, page, section_context(element), element.name, element.get_text(" ", strip=True))
            continue

        if element.name in MINOR_TAGS and not has_ancestor(element, GROUPING_TAGS):
            direct_parts = [
                clean_text(str(child))
                for child in element.children
                if isinstance(child, NavigableString) and not isinstance(child, Comment) and clean_text(str(child))
            ]
            for part in direct_parts:
                add_row(rows, page, section_context(element), element.name, part)

    return rows


def dedupe_rows(rows: list[dict[str, object]]) -> tuple[list[dict[str, object]], Counter]:
    seen_exact: set[tuple[str, str, str, str]] = set()
    seen_shared: dict[tuple[str, str, str], str] = {}
    omitted = Counter()
    output: list[dict[str, object]] = []

    for row in rows:
        text_fp = fingerprint(str(row["Text string"]))
        context = str(row["Approx. section/context"])
        source = str(row["Element / tag / attribute source"])
        page = str(row["Page filename"])
        exact_key = (page, context, source, text_fp)
        if exact_key in seen_exact:
            omitted["same-page exact duplicate"] += 1
            continue
        seen_exact.add(exact_key)

        if context in {"Navigation (shared)", "Footer (shared)"}:
            shared_key = (context, source, text_fp)
            if shared_key in seen_shared:
                omitted[f"repeated {context.lower()}"] += 1
                continue
            seen_shared[shared_key] = page
            row["Extraction note"] = "Shared sitewide copy; duplicate appearances omitted from later pages."

        output.append(row)

    return output, omitted


def build_workbook(rows: list[dict[str, object]], omitted: Counter) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Copy Review"

    notes = wb.create_sheet("README - Notes")
    summary = wb.create_sheet("Page Summary")

    title_fill = PatternFill("solid", fgColor="243447")
    header_fill = PatternFill("solid", fgColor="D8EAF5")
    light_fill = PatternFill("solid", fgColor="F7FAFC")
    note_fill = PatternFill("solid", fgColor="FFF6D6")
    white_font = Font(color="FFFFFF", bold=True)
    header_font = Font(color="1D2B36", bold=True)
    body_font = Font(color="1D2B36", size=10)
    thin = Side(style="thin", color="D9E2EA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for row in rows:
        ws.append([row[h] for h in HEADERS])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border

    widths = {
        "A": 26,
        "B": 44,
        "C": 26,
        "D": 72,
        "E": 16,
        "F": 24,
        "G": 18,
        "H": 42,
        "I": 16,
        "J": 34,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].alignment = Alignment(horizontal="right", vertical="top")

    for idx in range(2, ws.max_row + 1):
        if idx % 2 == 0:
            for cell in ws[idx]:
                cell.fill = light_fill

    if ws.max_row >= 2:
        table = Table(displayName="CopyReviewTable", ref=f"A1:J{ws.max_row}")
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        ws.add_table(table)

    status_validation = DataValidation(type="list", formula1='"Not started,In review,Approved,Needs rewrite,Question"', allow_blank=True)
    priority_validation = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(status_validation)
    ws.add_data_validation(priority_validation)
    status_validation.add(f"G2:G{max(ws.max_row, 2)}")
    priority_validation.add(f"I2:I{max(ws.max_row, 2)}")

    summary.sheet_view.showGridLines = False
    summary["A1"] = "FCG Website Copy Review - Page Summary"
    summary["A1"].fill = title_fill
    summary["A1"].font = white_font
    summary["A1"].alignment = Alignment(vertical="center")
    summary.merge_cells("A1:D1")
    summary["A3"] = "Page filename"
    summary["B3"] = "Extracted rows"
    summary["C3"] = "Approx. visible/UX strings included"
    summary["D3"] = "Review note"
    for cell in summary[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    counts = Counter(str(row["Page filename"]) for row in rows)
    for page in PAGES:
        summary.append([page, counts[page], counts[page], "Root live page included"])

    total_row = summary.max_row + 1
    summary[f"A{total_row}"] = "Total"
    summary[f"B{total_row}"] = len(rows)
    summary[f"C{total_row}"] = len(rows)
    summary[f"D{total_row}"] = "Rows in Copy Review sheet"

    for row in summary.iter_rows(min_row=4, max_row=summary.max_row):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in summary[total_row]:
        cell.font = Font(bold=True, color="1D2B36")
        cell.fill = note_fill
    for col, width in {"A": 30, "B": 16, "C": 34, "D": 42}.items():
        summary.column_dimensions[col].width = width

    notes.sheet_view.showGridLines = False
    notes["A1"] = "How to use this workbook"
    notes["A1"].fill = title_fill
    notes["A1"].font = white_font
    notes["A1"].alignment = Alignment(vertical="center")
    notes.merge_cells("A1:D1")

    note_rows = [
        ("Created", date.today().isoformat()),
        ("Scope", "12 live root HTML pages only. compare-v2-v3/** was intentionally excluded."),
        ("What is included", "Visible body copy plus useful title, meta description, alt, aria-label, and title attribute text."),
        ("What is excluded", "Scripts, CSS, SVG/path data, generic placeholder/logo alt text, code-like strings, and repeated shared nav/footer duplicates after first appearance."),
        ("How to review", "Use Suggested feedback owner, Status, Teammate comments, and Priority columns in the Copy Review sheet."),
        ("Status options", "Not started, In review, Approved, Needs rewrite, Question"),
        ("Priority options", "High, Medium, Low"),
    ]
    for r_idx, (label, value) in enumerate(note_rows, start=3):
        notes[f"A{r_idx}"] = label
        notes[f"B{r_idx}"] = value
        notes[f"A{r_idx}"].font = header_font
        notes[f"A{r_idx}"].fill = header_fill
        notes[f"B{r_idx}"].alignment = Alignment(wrap_text=True, vertical="top")
        notes[f"A{r_idx}"].border = border
        notes[f"B{r_idx}"].border = border

    omission_start = len(note_rows) + 5
    notes[f"A{omission_start}"] = "Deduplication / omissions"
    notes[f"A{omission_start}"].font = header_font
    notes[f"A{omission_start}"].fill = note_fill
    notes[f"A{omission_start}"].border = border
    notes[f"B{omission_start}"] = "Count"
    notes[f"B{omission_start}"].font = header_font
    notes[f"B{omission_start}"].fill = note_fill
    notes[f"B{omission_start}"].border = border
    for offset, (reason, count) in enumerate(sorted(omitted.items()), start=1):
        notes[f"A{omission_start + offset}"] = reason
        notes[f"B{omission_start + offset}"] = count
        notes[f"A{omission_start + offset}"].border = border
        notes[f"B{omission_start + offset}"].border = border

    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 100
    notes.column_dimensions["C"].width = 18
    notes.column_dimensions["D"].width = 18

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    alignment.vertical = alignment.vertical or "top"
                    cell.alignment = alignment

    wb.save(XLSX_PATH)


def write_csv(rows: list[dict[str, object]]) -> None:
    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    all_rows: list[dict[str, object]] = []
    for page in PAGES:
        all_rows.extend(extract_page(page))

    rows, omitted = dedupe_rows(all_rows)
    write_csv(rows)
    build_workbook(rows, omitted)
    print(f"xlsx={XLSX_PATH}")
    print(f"csv={CSV_PATH}")
    print(f"rows={len(rows)}")
    print("omitted=" + "; ".join(f"{reason}: {count}" for reason, count in sorted(omitted.items())))


if __name__ == "__main__":
    main()
